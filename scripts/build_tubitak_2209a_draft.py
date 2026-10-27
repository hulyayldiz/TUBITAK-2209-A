from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("output/docs/TUBITAK_2209A_GNSS_TYBS_Metin_Taslagi.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
TEAL = "1E6D78"
PALE = "EAF2F8"
PALE_TEAL = "E8F3F2"
GOLD = "FFF2CC"
GRAY = "5B6573"
WHITE = "FFFFFF"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.68)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string("252B33")
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for style_name, size, color in (("Title", 21, "000000"), ("Heading 1", 15, NAVY), ("Heading 2", 11, TEAL), ("Heading 3", 10, NAVY)):
    style = styles[style_name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True
title_ppr = styles["Title"]._element.pPr
if title_ppr is not None:
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, value, *, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(value))
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def table(headers, rows, widths=None, font_size=8.2):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.autofit = False
    set_repeat_table_header(tbl.rows[0])
    for j, label in enumerate(headers):
        set_cell_text(tbl.rows[0].cells[j], label, bold=True, color=WHITE, size=font_size)
        shade(tbl.rows[0].cells[j], NAVY)
    for i, values in enumerate(rows):
        cells = tbl.add_row().cells
        for j, value in enumerate(values):
            set_cell_text(cells[j], value, size=font_size)
            if i % 2 == 1:
                shade(cells[j], "F4F7FA")
    if widths:
        for row in tbl.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Inches(width)
    for row in tbl.rows:
        keep_row_together(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return tbl


def para(text="", *, style=None, bold_prefix=None, italic=False, color=None, size=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        a = p.add_run(bold_prefix)
        a.bold = True
        b = p.add_run(text[len(bold_prefix):])
        runs = [a, b]
    else:
        runs = [p.add_run(text)]
    for r in runs:
        r.italic = italic
        r.font.name = "Arial"
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.add_run(text)
    return p


def field(label, text, minimum=None):
    doc.add_heading(label, level=2)
    count = len(text)
    note = f" ({count:,} karakter" + (f"; asgari {minimum}" if minimum else "") + ")"
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("TYBS metni" + note)
    r.italic = True
    r.font.name = "Arial"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    p = para(text)
    p.paragraph_format.keep_together = True
    if minimum and count < minimum:
        raise ValueError(f"{label} has only {count} characters, needs {minimum}")
    return count


# Header and page number footer.
header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = header.add_run("2209-A  |  GNSS ARAŞTIRMA ÖNERİSİ TASLAĞI")
hr.font.name = "Arial"
hr.font.size = Pt(8)
hr.font.color.rgb = RGBColor.from_string(GRAY)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("TYBS alanlarına aktarım için düzenlenebilir çalışma metni  ·  Sayfa ")
fr.font.name = "Arial"
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor.from_string(GRAY)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

title = doc.add_paragraph(style="Title")
title.add_run("Yeni Bir GNSS Alıcısında Girişim Tespitinin Aktarılabilirliği")
subtitle = para("FGI Jammertest 2025 PolaRx5 kayıtlarında dondurulmuş değerlendirme", size=11, color=TEAL)
subtitle.paragraph_format.space_after = Pt(12)

note = para(
    "Kullanım notu: 2026 çağrısında öneri TYBS alanlarına girilir ve sistem tarafından oluşturulur; bu Word dosyası yükleme belgesi değildir. Sağlanan PDF 2024 döneminden ve başka bir projeye aittir. Başvuru sahibi ve danışman köşeli parantezli alanları tamamlamalıdır. Veri kataloğu ve küçük bir referans dosyası kontrol edilmiştir; saldırı senaryosu ölçümlerinin indirilmesi ve etiketlerinin doğrulanması henüz tamamlanmamıştır.",
    size=9,
)
note.paragraph_format.space_after = Pt(8)

doc.add_heading("Başvuru bilgileri", level=1)
table(
    ["Alan", "Taslak / tamamlanacak bilgi"],
    [
        ("Proje başlığı", "Yeni Bir GNSS Alıcısında Girişim Tespitinin Aktarılabilirliği"),
        ("Proje yürütücüsü", "[TYBS / ARBİS kayıtlarındaki güncel ad soyad]"),
        ("Proje ortakları", "[Varsa en fazla dört öğrenci; yoksa bireysel başvuru]"),
        ("Akademik danışman", "[Güncel YÖKSİS ve ARBİS kaydı doğrulanacak]"),
        ("Üniversite ve program", "[Yürütücünün kayıtlı olduğu üniversite, fakülte ve program]"),
        ("Bilimsel alan", "[TYBS sınıflandırmasından danışmanla birlikte seçilecek]"),
        ("Proje süresi", "12 ay; başlangıç tarihi destek kararı ve TYBS takvimi doğrultusunda belirlenecek"),
        ("Anahtar kelimeler", "GNSS; alıcılar arası aktarım; bağımsız değerlendirme; PolaRx5; spoofing; jamming"),
    ],
    widths=[1.55, 5.65],
)

doc.add_heading("Özet", level=1)
summary = (
    "Bu araştırma, mevcut GNSS ölçüm tutarlılığı izleme prototipinin daha önce geliştirme ve eşik ayarında kullanılmamış farklı bir alıcıya aktarılabilirliğini inceleyecektir. FGI Jammertest 2025 kataloğu Septentrio PolaRx5 ölçümleri, SBF/RINEX biçimleri ve kontrollü girişim senaryoları bildiriyor. 27 Ekim 2026 tarihli Metax manifestinde 102 dosya listeleniyor; katalog boyutları toplamı 874.580.271 bayttır; 27 Ekim 2026'da 299 baytlık statik referans koordinat dosyası indirildi, ancak aynı gün Metax dosya kullanılabilirlik ucu 405 `Download is not enabled` döndürdü; olay etiketleri dosya düzeyinde incelenmedi. Bu nedenle ilk hedef, araştırma için gerekli bir saldırı kaydı, eşleşen olay yer gerçeği ve kullanılabilir alanların erişimini doğrulamaktır. Erişim kapısı geçilirse dosya sürümü ve özeti kaydedilecek; test verisi açılmadan prototip kodu, profil, eşikler ve analiz protokolü dondurulacaktır. SBF/RINEX kayıtları deterministik adaptörle ortak şemaya aktarılacak; olmayan özellikler üretilmeyecektir. Jamming ve spoofing olayları için yeni alarm, kaçırılan olay ve gecikme hesaplanacak; yanlış alarm yalnızca açıkça temiz etiketlenmiş sürelerde raporlanacaktır. Mevcut TEXBAT ve Jammertest 2024 kayıtları geliştirme sürecinde incelendiği için bağımsız test sayılmayacaktır. Gerekli olay dosyalarına erişilemezse bağımsız performans sonucu iddia edilmeyecek ve başvuru kapsamı veri sağlayıcıyla teyit edilmiş uygulanabilir bir kaynağa göre yeniden belirlenmelidir."
)
words = len(summary.split())
if not 25 <= words <= 450:
    raise ValueError(f"Summary word count out of range: {words}")
p = para(summary)
p.paragraph_format.keep_together = True
p = para(f"Özet uzunluğu: {words} kelime (2026 çağrısı: 25-450 kelime).", size=8, color=GRAY, italic=True)

doc.add_heading("1 Araştırma önerisinin bilimsel niteliği", level=1)
scientific = (
    "GNSS spoofing, alıcının yanlış fakat makul görünen ölçümlere yönelmesine; jamming ise sinyal hizmetinin bozulmasına yol açabilir ve iki tehdit ayrı değerlendirilmelidir. Literatür, tespit performansının alıcıda erişilebilen sinyal/ölçüm özelliklerine ve alıcıya özgü kalibrasyona bağlı olduğunu; yanlış alarm sınırının açık tanımlanması gerektiğini gösterir [1, 2]. Code-minus-carrier ölçütleri yardımcı olabilir, ancak multipath ve iyonosferden de etkilendiğinden tek başına saldırı kanıtı değildir [3, 4]. Bu nedenle bir kayıt üzerinde ayarlanan prototipin farklı alıcı ve veri biçiminde aynı davranışı göstereceği varsayılamaz. Etsin kataloğu Jammertest 2025 kayıtlarını PolaRx5, SBF/RINEX ve senaryo yer gerçeğiyle tanımlar [7]. 27 Ekim 2026 Metax manifestinde 102 dosya listeleniyor; katalog boyutları toplamı 874.580.271 bayttır; erişim Open ve lisans CC BY 4.0 görünür; ancak şimdiye kadar indirilebilen küçük `ground_truth_static.csv` yalnızca statik referans koordinatlarını içerir. Dosya kullanılabilirlik ucu 405 `Download is not enabled` döndürdü; senaryo ölçümleri ve olay etiketleri edinilmedi. Bu nedenle önerinin bilimsel sorusu şu koşulla sınanacaktır: “Saldırı kayıtları ve olay yer gerçeği edinilebilirse, önceden dondurulmuş GNSS tutarsızlık izleyicisi daha önce kullanılmamış PolaRx5 kayıtlarında jamming/spoofing olaylarını hangi kapsam ve gecikmeyle belirler; açıkça temiz dönem varsa ne sıklıkta alarm üretir?” Mevcut TEXBAT ve Jammertest 2024 sonuçları geliştirme sürecinde incelendiğinden bağımsız test değildir. Çalışma evrensel tespit başarısı iddia etmeyecek; veri erişimini, gözlenebilir alanları, uyarlamayı ve olay düzeyindeki performans sınırlarını raporlayacaktır."
)
field("1.1 Konunun önemi ve bilimsel niteliği", scientific, 350)

objectives = (
    "Projenin amacı, geliştirme verilerinden ayrı PolaRx5 kayıtlarında mevcut GNSS tutarsızlık izleyicisinin sıfır-atış aktarımını değerlendirmek ve kapsam sınırlarını raporlamaktır. İlk hedef, gerekli saldırı SBF/RINEX dosyalarına ve eşleşen olay yer gerçeğine fiilî erişimi doğrulamak; erişilebilir her dosyanın katalog sürümünü, boyutunu ve SHA-256 özetini kaydetmektir. İkinci hedef, yalnız erişilebilir ve gerekli alanları taşıdığı doğrulanan dosyalar için deterministik adaptör ve kalite kontrolleri geliştirmektir. Üçüncü hedef, gerçek olay ölçümleri açılmadan prototip, mevcut profil, eşikler ve metrikleri dondurmaktır. Dördüncü hedef, erişim koşulu sağlanırsa jamming/spoofing olayları için yeni alarm, kaçırma ve gecikmeyi ayrı raporlamaktır. Yanlış alarm oranı yalnızca yer gerçeğinde açıkça temiz olduğu belirtilen sürelerde hesaplanacaktır. Veri erişim kapısı geçilemezse bağımsız başarı iddiası veya test sonucu sunulmayacaktır. Başarı ölçütü yüksek performans vaat etmek değil, bağımsız değerlendirme yapılabilirliğini doğrulayıp sonucu ve sınırları yeniden üretilebilir biçimde raporlamaktır."
)
field("1.2 Amaç ve hedefler", objectives, 350)

doc.add_heading("2 Yöntem", level=1)
method = (
    "Araştırma, geliştirme ve ayar süreçlerinde kullanılmamış Jammertest 2025 PolaRx5 kayıtlarında kayıt temelli sıfır-atış değerlendirmesi olarak tasarlanmıştır. 27 Ekim 2026 Metax manifesti 102 dosya ve toplam 874.580.271 bayt payload boyutu (yaklaşık 834,06 MiB) bildiriyor; katalog Open ve CC BY 4.0 olarak işaretli. 27 Ekim 2026'da 299 baytlık statik referans koordinat dosyası indirildi; saldırı senaryosu SBF/RINEX kayıtları ve olay etiketleri henüz edinilmedi. Bu nedenle ilk iş paketi, en az bir saldırı senaryosu ölçümü ile eşleşen olay yer gerçeğini indirip şema ve zaman uyumunu doğrulamaktır. Bu geçiş kapısı sağlanmazsa performans değerlendirmesi başlatılmaz ve başvuru kapsamı veri erişimi teyit edilen kaynakla danışman eşliğinde yeniden belirlenir. Veri edinilirse katalogdaki dosya özeti ve lisans kaydedilir; SBF/RINEX dönüştürücüsü yayınlanmış alan sözleşmesi ve yapay test örnekleriyle hazırlanır. Gerçek saldırı etiketleri açılmadan prototip kodu, model, profil, eşikler, yapılandırma, adaptör ve analiz protokolü SHA-256 ile dondurulur. Ölçümler ortak şemaya aktarılır; zaman birimleri, epoch atlamaları, geçersiz fix ve eksik alanlar korunur. Olmayan C/N0, pseudorange, taşıyıcı fazı veya temiz etiket türetilmez. Etiketler yalnızca çıkarım tamamlandıktan sonra uygulanır; olay başına yeni alarm, kaçırma ve ilk alarm gecikmesi raporlanır. Yanlış alarm/saat yalnızca açıkça temiz/normal etiketlenmiş zamanlarda hesaplanır."
)
field("Yöntem", method, 350)
para("Ön çalışma ve yöntem sınırı", style="Heading 2")
para(
    "Mevcut raporlardaki ön çalışma ölçümleri başvuruya bağımsız başarı iddiası olarak taşınmamıştır. Jammertest 2024 replay'inde 33 kayıtlı olayın 12'sinde saldırı başlangıcından sonra yeni alarm kenarı oluşmuş; bu olaylarda ortanca gecikme 428,3 saniye ölçülmüştür [8]. Veri sağlayıcının olay dışı aralıkları sertifikalı temiz olmadığından yanlış alarm oranı hesaplanmamıştır. TEXBAT güncel replay'inde DS2 kapsamı %99,77 ve ilk yeni alarm gecikmesi 0,90 saniye; DS7 kapsamı %98,49 ve gecikme 5,60 saniyedir [5, 6]. DS3'te alarm saldırı başlangıcında zaten etkin olduğundan yeni alarm gecikmesi ölçülememiştir. Bu kayıtların profili/parametreleri geliştirme sürecinde incelendiğinden sonuçlar aynı benchmark için keşifseldir ve bağımsız test değildir. Proje öncesinde bu raporlarla ölçüm ve kod sürümü danışmanla tekrar doğrulanacaktır."
)

doc.add_heading("3 Proje yönetimi", level=1)
doc.add_heading("3.1 Çalışma takvimi", level=2)
para("İş paketlerinin başarı ağırlıkları toplamı %100'dür. Aylar destek kararını izleyen proje başlangıcına göre planlanmıştır.", size=9)
table(
    ["İş adımı", "Süre", "Başarı ölçütü ve doğrulama", "Ağırlık"],
    [
        ("1. Veri erişim kapısı ve protokol", "Ay 1-2", "En az bir saldırı ölçümü ve eşleşen olay yer gerçeği fiilen edinilir; lisans, sürüm, dosya özeti ve zaman şeması doğrulanır. Erişim sağlanamazsa performans analizi başlamaz ve danışmanla kapsam yeniden belirlenir.", "%15"),
        ("2. SBF/RINEX adaptörü ve kalite kontrolü", "Ay 2-4", "Erişilebilir kaynak biçimi ortak gözlem şemasına aktarılır. Saat, birim, eksik alan ve fix/epoch kalite testleri geçer; dönüşüm kayıpları sayılır.", "%25"),
        ("3. Prototip ve analiz protokolünü kilitleme", "Ay 4-5", "Kod, model, mevcut alıcı profili, eşikler ve metrikler SHA-256 manifestiyle saldırı sonuçları açılmadan dondurulur; test verisinden ayar yapılmaz.", "%15"),
        ("4. Bağımsız olay değerlendirmesi", "Ay 6-9", "Jamming ve spoofing etiketleri çıkarım sonrası uygulanır. Olay düzeyinde yeni alarm, kaçırma ve gecikme raporlanır; yanlış alarm yalnız doğrulanmış temiz süre varsa hesaplanır.", "%30"),
        ("5. Tekrarlanabilir analiz ve rapor", "Ay 10-12", "Sürüm, yapılandırma, manifest ve sonuç tablolarıyla analizi yeniden üreten betikler tamamlanır; bulgular ve veri/alıcının sınırları raporlanır.", "%15"),
    ],
    widths=[1.45, 0.72, 4.38, 0.48],
    font_size=7.7,
)

doc.add_heading("3.2 Risk yönetimi", level=2)
table(
    ["Risk", "Olasılık / etki", "B planı"],
    [
        ("Saldırı ölçümlerine veya eşleşen olay yer gerçeğine erişilememesi", "Yüksek / yüksek. Katalog dosya listesini ve açık lisansı bildiriyor; şimdiye kadar yalnızca 299 baytlık statik referans dosyası indirildi.", "İlk iş paketi veri erişim kapısıdır. Saldırı SBF/RINEX dosyası ve olay yer gerçeği edinilip şema doğrulanmadan performans analizi başlamaz. Erişim sağlanamazsa bağımsız performans sonucu iddia edilmez; danışmanla kapsam, erişimi doğrulanmış bir veri kaynağına göre değiştirilir."),
        ("SBF/RINEX şemasının veya olay zaman çizelgesinin prototiple uyumsuz olması", "Orta / yüksek. Projede bu biçimler için adaptör yok.", "Önce tek bir kaynak biçimi seçilip alan eşlemesi test edilir. Saat, birim veya olay etiketi doğrulanamayan kayıtlar performans paydasına alınmaz; uyumlu bir kayıt yoksa sonuç sınırlılığı raporlanır."),
        ("Saldırı zaman çizelgesi veya kaynak saati belirsizliği", "Orta / yüksek. Olay takvimi ile alıcı saati farklı olabilir.", "Kaynak belgeleri karşılaştırılır; belirsizlik raporlanır. Güvenilir başlangıç/bitiş yoksa tespit gecikmesi hesaplanmaz."),
        ("C/N0, CMC, IMU veya temiz etiketlerin kaynakta bulunmaması", "Yüksek / orta. Gözlenebilirler veri kümesine göre değişir.", "Eksik alan üretilmez ve başka kaynaktan doldurulmaz. Yalnızca mevcut ölçüt alt kümesi değerlendirilir; metrik kapsamı her veri kaynağı için ayrı verilir."),
        ("Sabit yöntemin saldırıları kaçırması veya temiz kayıtta alarm üretmesi", "Orta / yüksek. Ön çalışma bu olasılığı göstermiştir.", "Sonuç gizlenmeden raporlanır. Kilitli test verisine bakarak eşik değiştirilmez; yeni parametre denenirse yeni sürüm ayrı ve yeni bağımsız kayıtla değerlendirilir."),
        ("Kaynak lisansı veya veri indirme erişiminin sınırlı olması", "Orta / orta.", "Yalnız izin verilen kayıtlar işlenir; dağıtımda verinin kendisi yerine kaynak bağlantısı, lisans ve SHA-256 manifesti paylaşılır. Erişim yoksa kapsam resmi ve yerelde doğrulanmış kaynaklara indirilir."),
    ],
    widths=[1.7, 2.2, 3.1],
    font_size=7.6,
)

doc.add_heading("3.3 Araştırma olanakları", level=2)
table(
    ["Olanak", "Projede kullanım", "Doğrulama"],
    [
        ("Mevcut GNSS analiz prototipi ve kaynak kodu", "Kaynak adaptörleri, alıcı profili, ölçüt hesapları ve tekrar çalıştırılabilir analiz.", "Öğrenci ve danışman mevcut sürümü, yazılım lisansını ve proje kullanım iznini doğrulamalı."),
        ("Bilgisayar ve yazılım ortamı", "Kayıt replay'i, tablo analizi ve sonuç üretimi; çalışma yerel bilgisayarda yürütülebilir.", "Bellek, disk alanı ve kurumsal erişim imkânları başvuru öncesi envanterden teyit edilmeli."),
        ("Yayımlanmış kayıtlar", "FGI Jammertest 2025 kataloğunda 102 dosya listelenir; statik referans koordinat dosyası indirildi. Saldırı senaryo kayıtları bağımsız test adayıdır ancak henüz edinilmedi.", "Saldırı SBF/RINEX kayıtları ve olay yer gerçeği için indirmenin çalıştığı, alanların ve saatlerin çözümlendiği başvuru öncesi doğrulanmalıdır; önceki TEXBAT/Jammertest verileri bağımsız test olarak kullanılmamalıdır."),
        ("Laboratuvar / GNSS donanımı", "Yeni normal oturum veya bağımsız alıcı kaydı gerekecekse kullanılabilir altyapı.", "[Danışman ve bölüm tarafından laboratuvar, alıcı, referans sensör ve güvenli ölçüm ortamı doğrulanacak.]"),
    ],
    widths=[1.6, 2.8, 2.6],
    font_size=7.8,
)

h = doc.add_heading("4 Araştırma önerisinin yaygın etkisi", level=1)
h.paragraph_format.page_break_before = True
para(
    "Proje, önceden geliştirme verilerinde incelenmiş GNSS tutarsızlık prototipinin bağımsız PolaRx5 ölçümlerine aktarılabilirliğini sınamak üzere tasarlanmıştır; bu sonuç, saldırı ölçümleri ve olay yer gerçeğine erişim doğrulanmasına bağlıdır. Beklenen çıktılar; (i) veri sürümü, lisansı, dosya özetleri ve gözlenebilir alan manifesti, (ii) edinilebilen biçim için deterministik adaptör ve kalite kontrolleri, (iii) erişim kapısı geçilirse jamming/spoofing olayları için dondurulmuş alarm, kaçırma ve gecikme raporu, (iv) tekrarlanabilir analiz betikleri ve öğrencinin GNSS veri işleme becerileridir. Paylaşım izni bulunan kod ve yönergeler yayımlanabilir; veri dosyaları lisans izin vermedikçe dağıtılmaz. Çalışma yüksek tespit başarısı veya sahada güvenli kullanım iddiası üretmez."
)

doc.add_heading("5 Belirtmek istediğiniz diğer konular", level=1)
para(
    "Bu çalışma canlı RF yayını, girişim üretimi, uçuş kontrolü veya insan katılımcı içermeyen çevrimdışı kayıt analizidir. GPS/konum kayıtları hassas olabileceğinden yalnızca kullanım izni doğrulanmış veri işlenecek; gereksiz koordinatlar yayımlanmayacak ve veri lisansına uyulacaktır. Etik kurul gerekliliği danışman ve yürütücü kurumla teyit edilmelidir. Prototip alarmı GNSS ölçüm tutarsızlığı göstergesidir; tek başına spoofing veya jamming nedenini ispatlamaz. Jammertest 2025 dosyalarına erişim ve biçim uyumu doğrulanmadan bağımsız performans sonucu vaat edilmez."
)

doc.add_heading("6 Kaynaklar", level=1)
refs = [
    "[1] M. L. Psiaki ve T. E. Humphreys, “GNSS Spoofing and Detection,” Proceedings of the IEEE, 104(6), 1258-1270, 2016. DOI: 10.1109/JPROC.2016.2526658.",
    "[2] F. Rothmaier, L. Taleghani, Y.-H. Chen, S. Lo, E. Phelts ve T. Walter, “GNSS Spoofing Detection through Metric Combinations: Calibration and Application of a General Framework,” Proceedings of ION GNSS+ 2021, 4249-4263. DOI: 10.33012/2021.18126.",
    "[3] A. M. Khan, N. Iqbal, A. A. Khan, M. F. Khan ve A. Ahmad, “Detection of Intermediate Spoofing Attack on Global Navigation Satellite System Receiver Through Slope Based Metrics,” The Journal of Navigation, 73(5), 1052-1068, 2020. DOI: 10.1017/S0373463320000168.",
    "[4] Q. Li ve diğerleri, “Intrinsic Identification and Mitigation of Multipath for Enhanced GNSS Positioning,” Sensors, 21(1), 188, 2021. DOI: 10.3390/s21010188.",
    "[5] T. E. Humphreys ve diğerleri, “The Texas Spoofing Test Battery: Toward a Standard for Evaluating GPS Signal Authentication Techniques,” Proceedings of ION GNSS 2012, 3569-3583. Veri ve senaryo açıklaması: UT Austin Radionavigation Laboratory, TEXBAT, https://radionavlab.ae.utexas.edu/texbat/.",
    "[6] M. Lemmenes ve diğerleri, “Detailed Analysis of the Texas Spoofing Test Battery Using a High Fidelity Software Receiver,” Proceedings of ION GNSS+ 2016. https://radionavlab.ae.utexas.edu/images/stories/files/papers/LemmenesGNSSpaper.pdf.",
    "[7] Finnish Geospatial Research Institute, “FGI Jammertest 2025 GNSS interference dataset,” Fairdata Etsin, 2026. Katalog 102 dosya ve toplam 874.580.271 bayt payload boyutu, Open erişim, CC BY 4.0 ve PolaRx5 SBF/RINEX kayıtlarını bildirir. 27 Ekim 2026'da indirilen `ground_truth_static.csv` statik referans koordinatları içerir; saldırı kayıtları/olay etiketleri indirilip doğrulanmamıştır. Katalog metadata manifesti `data_manifests/fgi_jammertest_2025.json` içindedir. https://etsin.fairdata.fi/dataset/095782a8-4e31-456e-8719-7021b4c96362.",
    "[8] Jammertest 2024, Zenodo sürüm 3, GNSS interference dataset. DOI: 10.5281/zenodo.15911589. Kullanılan olay zamanları ve kapsam için kaynak belgeleriyle tekrar doğrulama yapılacaktır.",
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.add_run(ref)

doc.add_heading("Başvuru öncesi tamamlanacak bilgiler", level=1)
for item in [
    "TYBS / ARBİS kayıtlarındaki öğrenci adı, ortaklar, program, danışman ve akademik alanı doğrulayın. Ek PDF’deki 2024 başvuru bilgilerini GNSS projesine otomatik aktarmayın.",
    "Danışmanla proje kapsamını ve 12 aylık takvimi netleştirin. Başvuru öncesi en az bir SBF/RINEX saldırı kaydının ve ona karşılık gelen olay yer gerçeğinin indirilmesini, şemasının ve zamanının çözümlenmesini doğrulayın.",
    "Katalogdaki `ground_truth_static.csv` yalnızca statik koordinat referansıdır; saldırı başlangıç/bitiş etiketi değildir. Gerçek olay etiketlerinin kaynağını, saat dilimini ve temiz süre tanımını ayrıca doğrulayın. Bu erişim kapısı geçilmezse bağımsız başarı vaadi vermeyin ve kapsamı danışmanla değiştirin.",
    "Bütçe ekranında yalnızca doğrudan gerekli, uygun ve fiyat teklifiyle gerekçelendirilebilen giderleri yazın. 2026 program sayfasındaki üst sınır 12.000 TL'dir; konferans, yayın/patent, konaklama ve yevmiye kalemlerini eklemeyin. Gerçek teklifleri danışman ve kurumla doğrulayın.",
    "TYBS önizlemesinde metin, tablo ve özel karakterlerin doğru aktarıldığını kontrol edin. Başvuru ve danışman onayını son tarihten önce tamamlayın.",
    "Üretken yapay zekâ bu taslağın araştırma önerisi metnini oluşturma ve düzenleme aşamasında kullanıldı. TYBS'deki ÜYZ beyanında bu kullanım dürüstçe açıklanmalı; metin, yöntem ve kaynakların doğruluğunu başvuru sahibi ile danışman kontrol etmelidir.",
]:
    bullet(item)

para(
    "2026 çağrı takvimi: başvuru ve danışman onayı 11 Kasım 2026 saat 17.30'a kadar; kuruluş yetkilisi e-imza aşaması 16-23 Kasım 2026. Güncel TYBS ekranı ve TÜBİTAK duyurusu başvuru anında yeniden kontrol edilmelidir.",
    italic=True,
    color=GRAY,
    size=8.5,
)

doc.core_properties.title = "Yeni Bir GNSS Alıcısında Girişim Tespitinin Aktarılabilirliği"
doc.core_properties.subject = "2026 TYBS alanları için düzenlenebilir metin taslağı"
doc.core_properties.author = ""
doc.save(OUT)
print(f"created={OUT.resolve()}")
print(f"abstract_words={words}")
print(f"field_chars={len(scientific)},{len(objectives)},{len(method)}")
